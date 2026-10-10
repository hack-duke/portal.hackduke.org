"""Convert the protected DuQuantum Google Forms workbook to import CSV.

The command prints aggregate shape information only. It never prints workbook
cells, email addresses, or other participant data.
"""

from __future__ import annotations

import argparse
import csv
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


SHEET_NAME = "Form Responses 1"


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    return parser


def csv_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%m/%d/%Y %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%m/%d/%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    if not args.input.is_file():
        print("Input workbook does not exist or is not a regular file")
        return 2

    workbook = load_workbook(args.input, read_only=True, data_only=False)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            print("Required response sheet is missing")
            return 2
        sheet = workbook[SHEET_NAME]
        rows: list[list[str]] = []
        for row in sheet.iter_rows():
            if any(cell.data_type == "f" for cell in row):
                print("Workbook contains formulas; refusing participant import")
                return 2
            rows.append([csv_value(cell.value) for cell in row])
    finally:
        workbook.close()

    if not rows or not rows[0]:
        print("Response sheet is empty")
        return 2
    width = len(rows[0])
    if any(len(row) != width for row in rows):
        print("Response sheet has inconsistent row widths")
        return 2

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8-sig", newline="") as destination:
        writer = csv.writer(destination, lineterminator="\n")
        writer.writerows(rows)

    print(f"Workbook rows converted: {len(rows) - 1}")
    print(f"Workbook columns converted: {width}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
