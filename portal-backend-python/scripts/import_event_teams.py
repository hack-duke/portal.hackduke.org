"""Validate and reconcile the authoritative DuQuantum final-team workbook.

This command is dry-run by default. It reads both ``Master List`` and ``Team
List`` and refuses to continue if their final assignments disagree. Participant
links use exact email first, then an exact full name that is unique in both the
eligible roster and workbook. Fuzzy matching is intentionally never used.

Reports and console output contain aggregate counts and row/slot references,
not participant names or email addresses.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence


EVENT_SLUG = "duquantum-2026"
SOURCE_KIND = "final_teams_xlsx"
DEFAULT_INPUT = (
    Path(__file__).resolve().parents[2]
    / "data"
    / "private"
    / EVENT_SLUG
    / "final_teams.xlsx"
)
MASTER_SHEET = "Master List"
TEAM_SHEET = "Team List"

MASTER_HEADERS = (
    "Team #",
    "Team Name",
    "Track",
    "Assigned Challenge",
    "Choice # Received",
    "Team Size",
    "Additional Members Requested (original team)",
    "Pref 1",
    "Pref 2",
    "Pref 3",
    "Pref 4",
    "Pref 5",
    "Pref 6",
    "Member Name",
    "Email",
    "Discord",
    "University",
    "Degree Level",
    "Original Form Team",
    "Team Source",
    "Flag",
)
TEAM_HEADERS = (
    "Team #",
    "Team Name",
    "Track",
    "Assigned Challenge",
    "Size",
    "Additional Members Requested",
    "Members",
    "1st Choice",
    "2nd Choice",
    "3rd Choice",
    "4th Choice",
    "5th Choice",
    "6th Choice",
    "Choice # Received",
    "Team Source",
    "Merged?",
    "Notes / Original Preferences",
)
TEAM_PREF_HEADERS = (
    "1st Choice",
    "2nd Choice",
    "3rd Choice",
    "4th Choice",
    "5th Choice",
    "6th Choice",
)
EMAIL_PATTERN = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")


class TeamImportValidationError(ValueError):
    """A workbook error whose message is safe to print without source values."""


@dataclass(frozen=True)
class ParsedTeamMember:
    row_number: int
    position: int
    display_name: str
    email: str
    normalized_email: str
    discord_username: str | None
    normalized_name: str
    first_name: str | None = None
    last_name: str | None = None


@dataclass(frozen=True)
class ParsedTeam:
    row_number: int
    team_number: str
    name: str
    normalized_name: str
    assigned_track: str
    assigned_challenge: str
    formation_source: str
    requested_additional_teammates: int | None
    additional_teammates_detail: str | None
    challenge_preferences: tuple[str, ...]
    members: tuple[ParsedTeamMember, ...]
    source_record_id: str
    track_preference: str | None = None
    submitted_at: None = None


@dataclass(frozen=True)
class TeamParseResult:
    source_sha256: str
    master_rows: int
    team_rows: int
    teams: tuple[ParsedTeam, ...]

    @property
    def member_rows(self) -> int:
        return sum(len(team.members) for team in self.teams)

    @property
    def raw_rows(self) -> int:
        return self.team_rows


@dataclass(frozen=True)
class ReconciledMember:
    team: ParsedTeam
    member: ParsedTeamMember
    registration_id: Any | None
    match_status: str
    match_method: str | None
    match_detail: str


@dataclass(frozen=True)
class ReconciliationResult:
    members: tuple[ReconciledMember, ...]

    @property
    def counts(self) -> Mapping[str, int]:
        return Counter(item.match_status for item in self.members)


@dataclass(frozen=True)
class DatabaseSyncResult:
    counts: Mapping[str, int]
    reconciliation: ReconciliationResult


def normalize_text(value: Any) -> str:
    decomposed = unicodedata.normalize("NFKD", str(value or "").strip().casefold())
    without_marks = "".join(
        character for character in decomposed if not unicodedata.combining(character)
    )
    return " ".join(
        "".join(
            character if character.isalnum() else " " for character in without_marks
        ).split()
    )


def normalize_email(value: Any) -> str:
    return str(value or "").strip().casefold()


def _text(value: Any, field: str, row_number: int) -> str:
    result = str(value or "").strip()
    if not result:
        raise TeamImportValidationError(f"{field} is blank at source row {row_number}")
    return result


def _team_number(value: Any, row_number: int) -> str:
    try:
        number = float(value)
    except (TypeError, ValueError):
        raise TeamImportValidationError(
            f"invalid team number at source row {row_number}"
        )
    if not number.is_integer() or number <= 0:
        raise TeamImportValidationError(
            f"invalid team number at source row {row_number}"
        )
    return str(int(number))


def _positive_size(value: Any, row_number: int) -> int:
    try:
        size = int(value)
    except (TypeError, ValueError):
        raise TeamImportValidationError(f"invalid team size at source row {row_number}")
    if not 1 <= size <= 4:
        raise TeamImportValidationError(f"invalid team size at source row {row_number}")
    return size


def _additional_teammates(value: Any) -> tuple[int | None, str | None]:
    raw = str(value or "").strip()
    if not raw:
        return None, None
    try:
        number = float(raw)
    except ValueError:
        return None, raw[:500]
    if number.is_integer() and 0 <= number <= 4:
        return int(number), None
    return None, raw[:500]


def _source_record_id(team_number: str, normalized_name: str) -> str:
    key = f"{EVENT_SLUG}\0final-team\0{team_number}\0{normalized_name}"
    return hashlib.sha256(key.encode("utf-8")).hexdigest()


def _validate_headers(
    actual: Sequence[str], required: Sequence[str], sheet: str
) -> None:
    missing = set(required).difference(actual)
    if missing:
        raise TeamImportValidationError(
            f"{sheet} is missing {len(missing)} required column(s)"
        )


def parse_final_rows(
    *,
    master_headers: Sequence[str],
    master_rows: Iterable[Mapping[str, Any]],
    team_headers: Sequence[str],
    team_rows: Iterable[Mapping[str, Any]],
    source_sha256: str,
) -> TeamParseResult:
    """Parse and cross-check the two final-assignment sheets."""
    _validate_headers(master_headers, MASTER_HEADERS, MASTER_SHEET)
    _validate_headers(team_headers, TEAM_HEADERS, TEAM_SHEET)

    master_groups: dict[str, list[tuple[int, Mapping[str, Any]]]] = defaultdict(list)
    seen_member_emails: set[str] = set()
    seen_member_names: set[str] = set()
    master_count = 0
    for row_number, row in enumerate(master_rows, start=2):
        master_count += 1
        number = _team_number(row.get("Team #"), row_number)
        name = _text(row.get("Member Name"), "member name", row_number)
        email = normalize_email(row.get("Email"))
        if not EMAIL_PATTERN.fullmatch(email):
            raise TeamImportValidationError(
                f"invalid member email at {MASTER_SHEET} row {row_number}"
            )
        normalized_name = normalize_text(name)
        if email in seen_member_emails:
            raise TeamImportValidationError("duplicate member email in Master List")
        if normalized_name in seen_member_names:
            raise TeamImportValidationError(
                "duplicate normalized member name in Master List"
            )
        seen_member_emails.add(email)
        seen_member_names.add(normalized_name)
        master_groups[number].append((row_number, row))

    parsed_teams: list[ParsedTeam] = []
    seen_numbers: set[str] = set()
    seen_team_names: set[str] = set()
    team_count = 0
    for row_number, row in enumerate(team_rows, start=2):
        team_count += 1
        number = _team_number(row.get("Team #"), row_number)
        if number in seen_numbers:
            raise TeamImportValidationError("duplicate team number in Team List")
        seen_numbers.add(number)
        name = _text(row.get("Team Name"), "team name", row_number)
        normalized_name = normalize_text(name)
        if normalized_name in seen_team_names:
            raise TeamImportValidationError(
                "duplicate normalized team name in Team List"
            )
        seen_team_names.add(normalized_name)
        track = _text(row.get("Track"), "assigned track", row_number)
        challenge = _text(
            row.get("Assigned Challenge"), "assigned challenge", row_number
        )
        formation_source = _text(row.get("Team Source"), "team source", row_number)
        declared_size = _positive_size(row.get("Size"), row_number)
        listed_names = [
            item.strip()
            for item in str(row.get("Members") or "").split(";")
            if item.strip()
        ]
        if len(listed_names) != declared_size:
            raise TeamImportValidationError(
                f"Team List member count mismatch at source row {row_number}"
            )
        preferences = tuple(
            str(row.get(header) or "").strip()
            for header in TEAM_PREF_HEADERS
            if str(row.get(header) or "").strip()
        )
        if len({value.casefold() for value in preferences}) != len(preferences):
            raise TeamImportValidationError(
                f"duplicate challenge preference at source row {row_number}"
            )

        source_members = master_groups.get(number, [])
        if len(source_members) != declared_size:
            raise TeamImportValidationError(
                f"Master List team size mismatch for team row {row_number}"
            )
        members: list[ParsedTeamMember] = []
        for position, (master_row_number, master_row) in enumerate(
            source_members, start=1
        ):
            cross_fields = (
                ("Team Name", name),
                ("Track", track),
                ("Assigned Challenge", challenge),
                ("Team Source", formation_source),
            )
            if any(
                str(master_row.get(field) or "").strip() != expected
                for field, expected in cross_fields
            ):
                raise TeamImportValidationError(
                    f"cross-sheet assignment mismatch for team row {row_number}"
                )
            if (
                _positive_size(master_row.get("Team Size"), master_row_number)
                != declared_size
            ):
                raise TeamImportValidationError(
                    f"cross-sheet size mismatch for team row {row_number}"
                )
            display_name = _text(
                master_row.get("Member Name"), "member name", master_row_number
            )
            email = normalize_email(master_row.get("Email"))
            discord = str(master_row.get("Discord") or "").strip() or None
            members.append(
                ParsedTeamMember(
                    row_number=master_row_number,
                    position=position,
                    display_name=display_name,
                    email=email,
                    normalized_email=email,
                    discord_username=discord,
                    normalized_name=normalize_text(display_name),
                )
            )
        if Counter(map(normalize_text, listed_names)) != Counter(
            member.normalized_name for member in members
        ):
            raise TeamImportValidationError(
                f"cross-sheet member mismatch for team row {row_number}"
            )
        requested, additional_detail = _additional_teammates(
            row.get("Additional Members Requested")
        )
        parsed_teams.append(
            ParsedTeam(
                row_number=row_number,
                team_number=number,
                name=name,
                normalized_name=normalized_name,
                assigned_track=track,
                assigned_challenge=challenge,
                formation_source=formation_source,
                requested_additional_teammates=requested,
                additional_teammates_detail=additional_detail,
                challenge_preferences=preferences,
                members=tuple(members),
                source_record_id=_source_record_id(number, normalized_name),
            )
        )

    if set(master_groups) != seen_numbers:
        raise TeamImportValidationError("Master List and Team List team sets differ")
    return TeamParseResult(
        source_sha256=source_sha256,
        master_rows=master_count,
        team_rows=team_count,
        teams=tuple(sorted(parsed_teams, key=lambda team: int(team.team_number))),
    )


def parse_final_teams_xlsx(path: Path) -> TeamParseResult:
    from openpyxl import load_workbook

    source_bytes = path.read_bytes()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        missing_sheets = {MASTER_SHEET, TEAM_SHEET}.difference(workbook.sheetnames)
        if missing_sheets:
            raise TeamImportValidationError(
                f"workbook is missing {len(missing_sheets)} required sheet(s)"
            )

        def sheet_rows(sheet_name: str):
            values = workbook[sheet_name].iter_rows(values_only=True)
            header_row = next(values, None)
            if header_row is None:
                raise TeamImportValidationError(f"{sheet_name} is empty")
            headers = tuple(str(value or "").strip() for value in header_row)
            rows = [
                dict(zip(headers, row))
                for row in values
                if any(value not in (None, "") for value in row)
            ]
            return headers, rows

        master_headers, master_rows = sheet_rows(MASTER_SHEET)
        team_headers, team_rows = sheet_rows(TEAM_SHEET)
        return parse_final_rows(
            master_headers=master_headers,
            master_rows=master_rows,
            team_headers=team_headers,
            team_rows=team_rows,
            source_sha256=hashlib.sha256(source_bytes).hexdigest(),
        )
    finally:
        workbook.close()


parse_xlsx = parse_final_teams_xlsx


def _registration_name_keys(registration: Any) -> set[str]:
    keys = {
        normalize_text(
            f"{getattr(registration, 'first_name', '') or ''} "
            f"{getattr(registration, 'last_name', '') or ''}"
        )
    }
    preferred_name = getattr(registration, "preferred_name", None)
    if preferred_name:
        keys.add(
            normalize_text(
                f"{preferred_name} {getattr(registration, 'last_name', '') or ''}"
            )
        )
    return {key for key in keys if key}


def reconcile_members(
    teams: Sequence[ParsedTeam], registrations: Sequence[Any]
) -> ReconciliationResult:
    """Link only deterministic identities; never perform fuzzy matching."""
    registrations_by_email: dict[str, list[Any]] = defaultdict(list)
    registrations_by_name: dict[str, dict[str, Any]] = defaultdict(dict)
    for registration in registrations:
        email = normalize_email(
            getattr(registration, "normalized_email", None)
            or getattr(registration, "email", None)
        )
        if email:
            registrations_by_email[email].append(registration)
        for key in _registration_name_keys(registration):
            registrations_by_name[key][str(registration.id)] = registration

    email_occurrences = Counter(
        member.normalized_email for team in teams for member in team.members
    )
    name_occurrences = Counter(
        member.normalized_name for team in teams for member in team.members
    )
    preliminary: list[ReconciledMember] = []
    for team in teams:
        for member in team.members:
            if email_occurrences[member.normalized_email] > 1:
                preliminary.append(
                    ReconciledMember(
                        team,
                        member,
                        None,
                        "ambiguous",
                        "email",
                        "email_listed_multiple_times_in_final_workbook",
                    )
                )
                continue
            email_matches = registrations_by_email.get(member.normalized_email, [])
            if len(email_matches) == 1:
                preliminary.append(
                    ReconciledMember(
                        team,
                        member,
                        email_matches[0].id,
                        "matched",
                        "email",
                        "exact_email",
                    )
                )
                continue
            if len(email_matches) > 1:
                preliminary.append(
                    ReconciledMember(
                        team,
                        member,
                        None,
                        "ambiguous",
                        "email",
                        "email_matches_multiple_registrations",
                    )
                )
                continue
            name_matches = list(
                registrations_by_name.get(member.normalized_name, {}).values()
            )
            if name_occurrences[member.normalized_name] == 1 and len(name_matches) == 1:
                preliminary.append(
                    ReconciledMember(
                        team,
                        member,
                        name_matches[0].id,
                        "matched",
                        "exact_unique_name_email_mismatch",
                        "exact_unique_name_email_mismatch",
                    )
                )
            elif len(name_matches) > 1:
                preliminary.append(
                    ReconciledMember(
                        team,
                        member,
                        None,
                        "ambiguous",
                        "name",
                        "name_matches_multiple_registrations",
                    )
                )
            else:
                preliminary.append(
                    ReconciledMember(
                        team, member, None, "unmatched", "email", "email_not_found"
                    )
                )

    candidate_indexes: dict[str, list[int]] = defaultdict(list)
    for index, item in enumerate(preliminary):
        if item.registration_id is not None:
            candidate_indexes[str(item.registration_id)].append(index)
    for indexes in candidate_indexes.values():
        if len(indexes) <= 1:
            continue
        confidence = {"email": 2, "exact_unique_name_email_mismatch": 1}
        best = max(
            confidence.get(preliminary[index].match_method, 0) for index in indexes
        )
        winners = [
            index
            for index in indexes
            if confidence.get(preliminary[index].match_method, 0) == best
        ]
        if len(winners) == 1:
            for index in indexes:
                if index != winners[0]:
                    preliminary[index] = replace(
                        preliminary[index],
                        registration_id=None,
                        match_status="unmatched",
                        match_detail="registration_already_linked_by_higher_confidence_match",
                    )
        else:
            for index in indexes:
                preliminary[index] = replace(
                    preliminary[index],
                    registration_id=None,
                    match_status="ambiguous",
                    match_detail="registration_listed_multiple_times",
                )
    return ReconciliationResult(tuple(preliminary))


def sync_database(
    result: TeamParseResult,
    *,
    commit: bool,
    expected_match_counts: Mapping[str, int | None] | None = None,
) -> DatabaseSyncResult:
    from sqlalchemy.orm import joinedload
    from db import get_local_session
    from models import Event, EventRegistration, EventTeam, EventTeamMember

    session = get_local_session()
    counts: Counter[str] = Counter()
    try:
        event = session.query(Event).filter(Event.slug == EVENT_SLUG).one_or_none()
        if event is None:
            raise RuntimeError(
                "DuQuantum event seed is missing; apply migrations first"
            )
        registrations = list(
            session.query(EventRegistration)
            .filter(
                EventRegistration.event_id == event.id,
                EventRegistration.admission_status == "accepted",
                EventRegistration.rsvp_status == "confirmed",
            )
            .all()
        )
        registration_by_id = {
            registration.id: registration for registration in registrations
        }
        reconciliation = reconcile_members(result.teams, registrations)
        expected_match_counts = expected_match_counts or {}
        mismatches = [
            status
            for status in ("matched", "unmatched", "ambiguous")
            if expected_match_counts.get(status) is not None
            and reconciliation.counts.get(status, 0) != expected_match_counts[status]
        ]
        if mismatches:
            raise TeamImportValidationError(
                f"database match-count guard failed for {len(mismatches)} aggregate value(s)"
            )

        reconciled_by_number: dict[str, list[ReconciledMember]] = defaultdict(list)
        for item in reconciliation.members:
            reconciled_by_number[item.team.team_number].append(item)
        existing_teams = list(
            session.query(EventTeam)
            .options(joinedload(EventTeam.members))
            .filter(EventTeam.event_id == event.id)
            .all()
        )
        existing_by_number = {
            team.team_number: team for team in existing_teams if team.team_number
        }
        existing_by_name = {team.normalized_name: team for team in existing_teams}
        selected_ids: set[Any] = set()
        current_database_teams: dict[str, Any] = {}
        for parsed in result.teams:
            by_number = existing_by_number.get(parsed.team_number)
            by_name = existing_by_name.get(parsed.normalized_name)
            if (
                by_number is not None
                and by_name is not None
                and by_number.id != by_name.id
            ):
                raise TeamImportValidationError("existing team identity conflict")
            database_team = by_number or by_name
            if database_team is None:
                database_team = EventTeam(
                    event_id=event.id,
                    name=parsed.name,
                    normalized_name=parsed.normalized_name,
                    source_record_id=parsed.source_record_id,
                )
                session.add(database_team)
                counts["teams_created"] += 1
            else:
                counts["teams_updated"] += 1
                if not database_team.active:
                    counts["teams_reactivated"] += 1
            database_team.name = parsed.name
            database_team.normalized_name = parsed.normalized_name
            database_team.team_number = parsed.team_number
            database_team.assigned_track = parsed.assigned_track
            database_team.assigned_challenge = parsed.assigned_challenge
            database_team.sponsor_name = parsed.assigned_challenge
            database_team.challenge_preferences = list(parsed.challenge_preferences)
            database_team.requested_additional_teammates = (
                parsed.requested_additional_teammates
            )
            database_team.additional_teammates_detail = (
                parsed.additional_teammates_detail
            )
            database_team.formation_source = parsed.formation_source
            database_team.source = SOURCE_KIND
            database_team.source_record_id = parsed.source_record_id
            database_team.source_submitted_at = None
            database_team.active = True
            current_database_teams[parsed.team_number] = database_team
            session.flush()
            selected_ids.add(database_team.id)

        for existing in existing_teams:
            if existing.id in selected_ids:
                continue
            if existing.active:
                existing.active = False
                counts["teams_archived"] += 1
            for member in existing.members:
                member.registration_id = None
                member.match_status = "unmatched"
                member.match_method = None
                member.match_detail = "team_not_in_latest_import"

        session.flush()
        for team in current_database_teams.values():
            for member in list(team.members):
                session.delete(member)
        session.flush()
        for parsed in result.teams:
            database_team = current_database_teams[parsed.team_number]
            for item in sorted(
                reconciled_by_number[parsed.team_number],
                key=lambda value: value.member.position,
            ):
                registration = registration_by_id.get(item.registration_id)
                session.add(
                    EventTeamMember(
                        event_id=event.id,
                        team_id=database_team.id,
                        registration_id=item.registration_id,
                        member_position=item.member.position,
                        display_name=item.member.display_name,
                        first_name=registration.first_name if registration else None,
                        last_name=registration.last_name if registration else None,
                        email=item.member.email,
                        normalized_email=item.member.normalized_email,
                        discord_username=item.member.discord_username,
                        match_status=item.match_status,
                        match_method=item.match_method,
                        match_detail=item.match_detail,
                        source_row_number=item.member.row_number,
                    )
                )
                counts[f"members_{item.match_status}"] += 1
        session.flush()
        if commit:
            session.commit()
        else:
            session.rollback()
        return DatabaseSyncResult(dict(counts), reconciliation)
    except Exception:
        session.rollback()
        raise
    finally:
        session.close()


def write_report(
    path: Path,
    *,
    parse_result: TeamParseResult,
    database_result: DatabaseSyncResult | None,
) -> None:
    reconciliation = database_result.reconciliation if database_result else None
    issues = (
        []
        if reconciliation is None
        else [
            {
                "team_number": item.team.team_number,
                "master_row_number": item.member.row_number,
                "member_position": item.member.position,
                "status": item.match_status,
                "method": item.match_method,
                "reason": item.match_detail,
            }
            for item in reconciliation.members
            if item.match_status != "matched"
        ]
    )
    payload = {
        "event_slug": EVENT_SLUG,
        "source_sha256": parse_result.source_sha256,
        "master_rows": parse_result.master_rows,
        "team_rows": parse_result.team_rows,
        "teams": len(parse_result.teams),
        "member_rows": parse_result.member_rows,
        "team_sizes": dict(
            sorted(Counter(len(team.members) for team in parse_result.teams).items())
        ),
        "tracks": dict(
            sorted(Counter(team.assigned_track for team in parse_result.teams).items())
        ),
        "formation_sources": dict(
            sorted(
                Counter(team.formation_source for team in parse_result.teams).items()
            )
        ),
        "match_counts": dict(reconciliation.counts) if reconciliation else None,
        "issues": issues,
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, sort_keys=True), encoding="utf-8")


def _print_summary(
    result: TeamParseResult, database_result: DatabaseSyncResult | None
) -> None:
    print(f"Master List rows: {result.master_rows}")
    print(f"Team List rows: {result.team_rows}")
    print(f"Final teams: {len(result.teams)}")
    print(f"Final member rows: {result.member_rows}")
    if database_result:
        matches = database_result.reconciliation.counts
        for status in ("matched", "unmatched", "ambiguous"):
            print(f"Members {status}: {matches.get(status, 0)}")
        for key, value in sorted(database_result.counts.items()):
            print(f"Database {key.replace('_', ' ')}: {value}")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--database-check", action="store_true")
    parser.add_argument("--commit", action="store_true")
    parser.add_argument("--confirm-event")
    parser.add_argument("--report", type=Path)
    parser.add_argument("--expect-source-sha256")
    parser.add_argument("--expect-master-rows", type=int)
    parser.add_argument("--expect-team-rows", type=int)
    parser.add_argument("--expect-teams", type=int)
    parser.add_argument("--expect-members", type=int)
    parser.add_argument("--expect-matched", type=int)
    parser.add_argument("--expect-unmatched", type=int)
    parser.add_argument("--expect-ambiguous", type=int)
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    required_commit_guards = (
        args.expect_source_sha256,
        args.expect_master_rows,
        args.expect_team_rows,
        args.expect_teams,
        args.expect_members,
        args.expect_matched,
        args.expect_unmatched,
        args.expect_ambiguous,
    )
    if args.commit and args.confirm_event != EVENT_SLUG:
        print(f"Refusing commit: pass --confirm-event {EVENT_SLUG}")
        return 2
    if args.commit and any(value is None for value in required_commit_guards):
        print(
            "Refusing commit: source, sheet, team, member, and match guards are required"
        )
        return 2
    if not args.input.is_file():
        print("Input XLSX does not exist or is not a regular file")
        return 2
    try:
        result = parse_xlsx(args.input)
    except (OSError, TeamImportValidationError, ValueError) as exc:
        print(f"Validation failed: {exc}")
        return 2
    expected = {
        "source": (args.expect_source_sha256, result.source_sha256),
        "master rows": (args.expect_master_rows, result.master_rows),
        "team rows": (args.expect_team_rows, result.team_rows),
        "teams": (args.expect_teams, len(result.teams)),
        "members": (args.expect_members, result.member_rows),
    }
    mismatches = [
        label
        for label, (wanted, actual) in expected.items()
        if wanted is not None and wanted != actual
    ]
    if mismatches:
        _print_summary(result, None)
        print(f"Source/count guard failed for {len(mismatches)} aggregate value(s)")
        return 2
    database_result = None
    if args.database_check or args.commit:
        try:
            database_result = sync_database(
                result,
                commit=args.commit,
                expected_match_counts={
                    "matched": args.expect_matched,
                    "unmatched": args.expect_unmatched,
                    "ambiguous": args.expect_ambiguous,
                },
            )
        except TeamImportValidationError as exc:
            print(f"Database validation failed: {exc}")
            return 2
        except Exception as exc:
            print(
                f"Database operation failed ({type(exc).__name__}); transaction rolled back"
            )
            return 1
    if args.report:
        write_report(args.report, parse_result=result, database_result=database_result)
    _print_summary(result, database_result)
    print(
        "Import committed" if args.commit else "Dry run only; no database changes made"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
